"""Чтение и запись pdf, docx, xlsx, xls. Тяжёлые библиотеки импортируются внутри функций."""

from __future__ import annotations

import json
import os
import zipfile
from collections.abc import Callable
from datetime import date, datetime
from pathlib import Path

from coding_mcp.errors import DocumentError
from coding_mcp.jail import open_for_read, prepare_write, relative_name
from coding_mcp.limits import (
    MAX_CELL_CHARS,
    MAX_COLS,
    MAX_INPUT_CHARS,
    MAX_PDF_PAGES,
    MAX_ROWS,
    MAX_SHEETS,
    MAX_UNCOMPRESSED_BYTES,
)
from coding_mcp.runtime import clip

_ILLEGAL_SHEET = set(":\\/?*[]")


def read_pdf(root: Path, raw: str) -> str:
    path = open_for_read(root, raw, ".pdf")
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise DocumentError("PDF зашифрован.")
        total = len(reader.pages)
        chunks: list[str] = []
        limit = min(total, MAX_PDF_PAGES)
        for index in range(limit):
            extracted = reader.pages[index].extract_text() or ""
            body = extracted.strip() or "[нет текстового слоя]"
            chunks.append(f"## page {index + 1}\n{body}")
        if total > limit:
            chunks.append(f"[пропущены страницы {limit + 1}–{total}]")
    except DocumentError:
        raise
    except PdfReadError:
        raise DocumentError("Не удалось прочитать PDF.") from None
    text = "\n\n".join(chunks).strip()
    if not text:
        raise DocumentError("Текст в PDF не найден.")
    return clip(text)


def read_docx(root: Path, raw: str) -> str:
    path = open_for_read(root, raw, ".docx")
    _reject_zip_bomb(path)
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(str(path))
    except (PackageNotFoundError, zipfile.BadZipFile):
        raise DocumentError("Не удалось прочитать DOCX.") from None

    lines: list[str] = []
    rows_seen = 0
    truncated = False
    for block in _iter_blocks(document, Paragraph, Table):
        if isinstance(block, Paragraph):
            line = _paragraph_line(block)
            if line:
                lines.append(line)
            continue
        table_lines, rows_seen, table_truncated = _table_lines(block.rows, rows_seen)
        lines.extend(table_lines)
        truncated = truncated or table_truncated
        if rows_seen >= MAX_ROWS:
            truncated = True
            break
    if truncated:
        lines.append("[часть таблиц пропущена по лимиту строк]")
    text = "\n".join(lines).strip()
    if not text:
        raise DocumentError("Текст в DOCX не найден.")
    return clip(text)


def write_docx(root: Path, raw: str, text: str) -> str:
    if not isinstance(text, str) or not text.strip():
        raise DocumentError("Текст пустой.")
    if len(text) > MAX_INPUT_CHARS:
        raise DocumentError("Текст превышает лимит размера.")
    blocks = _docx_blocks(text)
    path = prepare_write(root, raw, ".docx")

    def save(target: Path) -> None:
        from docx import Document

        document = Document()
        for kind, payload in blocks:
            if kind == "heading":
                level, line = payload
                document.add_heading(line, level=level)
            else:
                document.add_paragraph(payload)
        document.save(str(target))

    _atomic_save(path, save)
    return f"Записан файл {relative_name(root, path)}"


def read_xlsx(root: Path, raw: str, sheet: str | None = None) -> str:
    path = open_for_read(root, raw, ".xlsx")
    _reject_zip_bomb(path)
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile):
        raise DocumentError("Не удалось прочитать XLSX.") from None
    try:
        return _render_openpyxl(workbook, sheet)
    finally:
        workbook.close()


def write_xlsx(root: Path, raw: str, content: dict | str) -> str:
    sheets = _parse_sheets(content)
    path = prepare_write(root, raw, ".xlsx")

    def save(target: Path) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)
        for name, rows in sheets:
            worksheet = workbook.create_sheet(name)
            for row in rows:
                worksheet.append(row)
        workbook.save(str(target))

    _atomic_save(path, save)
    return f"Записан файл {relative_name(root, path)}"


def read_xls(root: Path, raw: str, sheet: str | None = None) -> str:
    path = open_for_read(root, raw, ".xls")
    import xlrd

    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except xlrd.XLRDError as exc:
        message = str(exc).lower()
        if "encrypt" in message:
            raise DocumentError("XLS зашифрован.") from None
        raise DocumentError("Не удалось прочитать XLS.") from None
    try:
        return _render_xlrd(book, sheet)
    finally:
        book.release_resources()


def _render_openpyxl(workbook, sheet: str | None) -> str:
    names = list(workbook.sheetnames)
    selected, note = _select_sheets(names, sheet)
    chunks: list[str] = []
    for name in selected:
        worksheet = workbook[name]
        rows: list[list[str]] = []
        extra = False
        for index, row in enumerate(
            worksheet.iter_rows(max_col=MAX_COLS, values_only=True)
        ):
            if index >= MAX_ROWS:
                extra = True
                break
            rows.append([_cell_text(value) for value in row])
        body = _tsv(rows)
        if extra:
            body = f"{body}\n[строки после {MAX_ROWS} пропущены]" if body else f"[строки после {MAX_ROWS} пропущены]"
        chunks.append(f"## {name}\n{body}".rstrip())
    if note:
        chunks.append(note)
    return clip("\n\n".join(chunks))


def _render_xlrd(book, sheet: str | None) -> str:
    names = book.sheet_names()
    selected, note = _select_sheets(names, sheet)
    chunks: list[str] = []
    for name in selected:
        worksheet = book.sheet_by_name(name)
        limit = min(worksheet.nrows, MAX_ROWS)
        rows: list[list[str]] = []
        for row_index in range(limit):
            values: list[str] = []
            cols = min(worksheet.ncols, MAX_COLS)
            for col_index in range(cols):
                cell = worksheet.cell(row_index, col_index)
                values.append(_xlrd_text(book, cell))
            rows.append(values)
        body = _tsv(rows)
        if worksheet.nrows > MAX_ROWS:
            suffix = f"[строки после {MAX_ROWS} пропущены]"
            body = f"{body}\n{suffix}" if body else suffix
        chunks.append(f"## {name}\n{body}".rstrip())
    if note:
        chunks.append(note)
    return clip("\n\n".join(chunks))


def _select_sheets(names: list[str], sheet: str | None) -> tuple[list[str], str]:
    if not names:
        raise DocumentError("В книге нет листов.")
    if sheet is not None:
        if not isinstance(sheet, str) or not sheet.strip():
            raise DocumentError("Имя листа пустое.")
        if sheet not in names:
            raise DocumentError("Лист не найден.")
        return [sheet], ""
    selected = names[:MAX_SHEETS]
    note = ""
    if len(names) > MAX_SHEETS:
        note = f"[листы после первых {MAX_SHEETS} пропущены]"
    return selected, note


def _parse_sheets(content: dict | str) -> list[tuple[str, list[list[object]]]]:
    if isinstance(content, str):
        if len(content) > MAX_INPUT_CHARS:
            raise DocumentError("Запрос превышает лимит размера.")
        try:
            parsed = json.loads(content)
        except json.JSONDecodeError:
            raise DocumentError("content должен быть объектом с полем sheets.") from None
    elif isinstance(content, dict):
        parsed = content
    else:
        raise DocumentError("content должен быть объектом с полем sheets.")

    sheets = parsed.get("sheets") if isinstance(parsed, dict) else None
    if not isinstance(sheets, list) or not sheets:
        raise DocumentError("Нужен непустой список sheets.")
    if len(sheets) > MAX_SHEETS:
        raise DocumentError(f"Слишком много листов, максимум {MAX_SHEETS}.")

    ready: list[tuple[str, list[list[object]]]] = []
    seen: set[str] = set()
    for sheet in sheets:
        if not isinstance(sheet, dict):
            raise DocumentError("Описание листа должно быть объектом.")
        name = _sheet_name(sheet.get("name"))
        if name in seen:
            raise DocumentError("Имена листов повторяются.")
        seen.add(name)
        rows = sheet.get("rows")
        if not isinstance(rows, list):
            raise DocumentError("У листа нет списка rows.")
        if len(rows) > MAX_ROWS:
            raise DocumentError(f"Слишком много строк, максимум {MAX_ROWS}.")
        cleaned: list[list[object]] = []
        for row in rows:
            if not isinstance(row, list):
                raise DocumentError("Строка листа должна быть списком ячеек.")
            if len(row) > MAX_COLS:
                raise DocumentError(f"Слишком много столбцов, максимум {MAX_COLS}.")
            cleaned.append([_write_cell(value) for value in row])
        ready.append((name, cleaned))
    return ready


def _sheet_name(name: object) -> str:
    if not isinstance(name, str):
        raise DocumentError("У листа нет имени.")
    cleaned = name.strip()
    if not cleaned or len(cleaned) > 31 or any(char in _ILLEGAL_SHEET for char in cleaned):
        raise DocumentError("Недопустимое имя листа.")
    return cleaned


def _write_cell(value: object) -> object:
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int) and not isinstance(value, bool):
        return value
    if isinstance(value, float):
        if value != value or value in {float("inf"), float("-inf")}:
            raise DocumentError("Недопустимое число в ячейке.")
        return value
    if isinstance(value, str):
        if len(value) > MAX_CELL_CHARS:
            raise DocumentError("Слишком длинное значение ячейки.")
        return _strip_controls(value)
    raise DocumentError("Ячейка должна быть строкой, числом, bool или null.")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ")
    text = _strip_controls(str(value)).replace("\t", " ").replace("\n", " ")
    if len(text) > MAX_CELL_CHARS:
        return text[:MAX_CELL_CHARS] + "…"
    return text


def _xlrd_text(book, cell) -> str:
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat(sep=" ")
        except xlrd.XLDateError:
            return ""
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "true" if cell.value else "false"
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
        return ""
    return _cell_text(cell.value)


def _tsv(rows: list[list[str]]) -> str:
    while rows and not any(cell.strip() for cell in rows[-1]):
        rows.pop()
    return "\n".join("\t".join(row).rstrip("\t") for row in rows)


def _docx_blocks(text: str) -> list[tuple[str, object]]:
    blocks: list[tuple[str, object]] = []
    for raw_line in text.splitlines():
        line = _strip_controls(raw_line).strip()
        if not line:
            continue
        level = _heading_level(line)
        if level:
            heading = line[level + 1 :].strip()
            if not heading:
                raise DocumentError("Пустой заголовок.")
            blocks.append(("heading", (level, heading)))
        else:
            blocks.append(("paragraph", line))
    if not blocks:
        raise DocumentError("Текст пустой.")
    return blocks


def _heading_level(line: str) -> int:
    if not line.startswith("#"):
        return 0
    marks = 0
    for char in line:
        if char != "#":
            break
        marks += 1
    if marks > 3 or len(line) <= marks or line[marks] != " ":
        return 0
    return marks


def _paragraph_line(paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    style_name = ""
    if paragraph.style is not None and paragraph.style.name:
        style_name = paragraph.style.name
    if style_name.startswith("Heading"):
        digits = [char for char in style_name if char.isdigit()]
        level = int(digits[0]) if digits else 1
        level = min(max(level, 1), 3)
        return f"{'#' * level} {_strip_controls(text)}"
    return _strip_controls(text)


def _table_lines(rows, rows_seen: int) -> tuple[list[str], int, bool]:
    lines: list[str] = []
    truncated = False
    for row in rows:
        if rows_seen >= MAX_ROWS:
            truncated = True
            break
        cells = [_strip_controls(cell.text).replace("\n", " ").strip() for cell in row.cells]
        lines.append("\t".join(cells))
        rows_seen += 1
    if lines:
        lines.append("")
    return lines, rows_seen, truncated


def _iter_blocks(document, paragraph_type, table_type):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield paragraph_type(child, document)
        elif child.tag.endswith("}tbl"):
            yield table_type(child, document)


def _reject_zip_bomb(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        raise DocumentError("Файл повреждён.") from None
    if total > MAX_UNCOMPRESSED_BYTES:
        raise DocumentError("Файл превышает лимит распакованного размера.")


def _atomic_save(path: Path, save: Callable[[Path], None]) -> None:
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        save(temporary)
        os.replace(temporary, path)
    except DocumentError:
        temporary.unlink(missing_ok=True)
        raise
    except PermissionError:
        temporary.unlink(missing_ok=True)
        raise DocumentError("Нет доступа к файлу.") from None
    except OSError:
        temporary.unlink(missing_ok=True)
        raise DocumentError("Не удалось записать файл.") from None
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _strip_controls(value: str) -> str:
    return "".join(char for char in value if char in "\n\t" or ord(char) >= 32)
