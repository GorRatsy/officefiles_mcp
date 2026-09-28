from pathlib import Path

import pytest

from coding_mcp.documents import read_docx, read_pdf, read_xls, read_xlsx, write_docx, write_xlsx
from coding_mcp.errors import DocumentError
from coding_mcp.limits import MAX_ROWS


def _pdf_bytes(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 18 Tf 72 100 Td ({escaped}) Tj ET\n".encode("ascii")
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 200] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n",
        f"4 0 obj\n<< /Length {len(stream)} >>\nstream\n".encode("ascii") + stream + b"endstream\nendobj\n",
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
    ]
    parts = [b"%PDF-1.4\n"]
    offsets: list[int] = []
    for obj in objects:
        offsets.append(sum(len(part) for part in parts))
        parts.append(obj)
    xref_at = sum(len(part) for part in parts)
    xref_lines = ["xref\n0 6\n", "0000000000 65535 f \n"]
    xref_lines.extend(f"{offset:010d} 00000 n \n" for offset in offsets)
    trailer = f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    return b"".join(parts) + "".join(xref_lines).encode("ascii") + trailer.encode("ascii")


def test_read_pdf(tmp_path: Path) -> None:
    root = tmp_path / "work"
    root.mkdir()
    (root / "hello.pdf").write_bytes(_pdf_bytes("Hello PDF"))

    text = read_pdf(root, "hello.pdf")

    assert "Hello PDF" in text
    assert "## page 1" in text


def test_pdf_outside_root_is_rejected(tmp_path: Path) -> None:
    root = tmp_path / "work"
    root.mkdir()
    outside = tmp_path / "hello.pdf"
    outside.write_bytes(_pdf_bytes("Secret"))

    with pytest.raises(DocumentError, match="вне рабочей директории"):
        read_pdf(root, str(outside))


def test_docx_roundtrip_with_table(tmp_path: Path) -> None:
    root = tmp_path / "work"
    root.mkdir()

    written = write_docx(root, "out/note.docx", "# Отчёт\nПервая строка\n## Раздел")
    assert written == "Записан файл out/note.docx"

    from docx import Document

    document = Document(str(root / "out" / "note.docx"))
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Имя"
    table.cell(0, 1).text = "Сумма"
    table.cell(1, 0).text = "Анна"
    table.cell(1, 1).text = "10"
    document.save(str(root / "out" / "note.docx"))

    text = read_docx(root, "out/note.docx")

    assert "# Отчёт" in text
    assert "Первая строка" in text
    assert "## Раздел" in text
    assert "Анна\t10" in text


def test_xlsx_roundtrip_keeps_types_and_formula(tmp_path: Path) -> None:
    root = tmp_path / "work"
    root.mkdir()
    content = {
        "sheets": [
            {
                "name": "Лист1",
                "rows": [["Имя", "Сумма", "Флаг"], ["Анна", 2, True], ["Борис", 2.5, None], [None, "=SUM(B2:B3)", False]],
            }
        ]
    }

    assert write_xlsx(root, "table.xlsx", content) == "Записан файл table.xlsx"
    text = read_xlsx(root, "table.xlsx", sheet="Лист1")

    assert "## Лист1" in text
    assert "Анна\t2\ttrue" in text
    assert "Борис\t2.5" in text
    assert "=SUM(B2:B3)\tfalse" in text


def test_write_xlsx_accepts_json_string(tmp_path: Path) -> None:
    root = tmp_path / "work"
    root.mkdir()

    write_xlsx(root, "one.xlsx", '{"sheets":[{"name":"S","rows":[["a",1]]}]}')

    assert "a\t1" in read_xlsx(root, "one.xlsx")


def test_write_rejects_too_many_rows(tmp_path: Path) -> None:
    root = tmp_path / "work"
    root.mkdir()
    rows = [["x"]] * (MAX_ROWS + 1)

    with pytest.raises(DocumentError, match="Слишком много строк"):
        write_xlsx(root, "big.xlsx", {"sheets": [{"name": "S", "rows": rows}]})
    assert not (root / "big.xlsx").exists()


def test_wrong_extension_is_rejected(tmp_path: Path) -> None:
    root = tmp_path / "work"
    root.mkdir()
    (root / "note.doc").write_text("no", encoding="utf-8")

    with pytest.raises(DocumentError, match="docx"):
        read_docx(root, "note.doc")


def test_xls_read(tmp_path: Path) -> None:
    import xlwt

    root = tmp_path / "work"
    root.mkdir()
    book = xlwt.Workbook()
    sheet = book.add_sheet("Данные")
    sheet.write(0, 0, "Имя")
    sheet.write(0, 1, "Сумма")
    sheet.write(1, 0, "Анна")
    sheet.write(1, 1, 7)
    book.save(str(root / "old.xls"))

    text = read_xls(root, "old.xls", sheet="Данные")

    assert "Имя\tСумма" in text
    assert "Анна\t7" in text


def test_missing_xls_sheet(tmp_path: Path) -> None:
    import xlwt

    root = tmp_path / "work"
    root.mkdir()
    book = xlwt.Workbook()
    book.add_sheet("Данные")
    book.save(str(root / "old.xls"))

    with pytest.raises(DocumentError, match="Лист не найден"):
        read_xls(root, "old.xls", sheet="Другой")


def test_output_is_clipped(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    root = tmp_path / "work"
    root.mkdir()
    rows = [["x" * 100] for _ in range(50)]
    write_xlsx(root, "wide.xlsx", {"sheets": [{"name": "S", "rows": rows}]})
    monkeypatch.setattr("coding_mcp.runtime.MAX_OUTPUT_CHARS", 120)

    text = read_xlsx(root, "wide.xlsx")

    assert len(text) <= 120 + len("\n\n[вывод обрезан по лимиту]")
    assert text.endswith("[вывод обрезан по лимиту]")


def test_zip_bomb_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    root = tmp_path / "work"
    root.mkdir()
    write_xlsx(root, "ok.xlsx", {"sheets": [{"name": "S", "rows": [["a"]]}]})
    monkeypatch.setattr("coding_mcp.documents.MAX_UNCOMPRESSED_BYTES", 10)

    with pytest.raises(DocumentError, match="распакованного размера"):
        read_xlsx(root, "ok.xlsx")


def test_call_timeout(monkeypatch: pytest.MonkeyPatch) -> None:
    from coding_mcp.runtime import run_bounded

    monkeypatch.setattr("coding_mcp.runtime.CALL_TIMEOUT_SECONDS", 0.05)

    def slow() -> None:
        import time

        time.sleep(2)

    with pytest.raises(DocumentError, match="время обработки"):
        run_bounded(slow, timeout=0.05)
